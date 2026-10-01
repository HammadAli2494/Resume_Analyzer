import io
import unittest
import zipfile

from docx import Document

from app.services.parse import extract_text, validate_upload


def _pdf(text: str) -> bytes:
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 18 Tf 72 720 Td ({safe}) Tj ET".encode("latin-1", "replace")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Count 1 /Kids [3 0 R] >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode())
        output.extend(body)
        output.extend(b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return bytes(output)


def _docx_bytes(builder) -> bytes:
    document = Document()
    builder(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


LONG = (
    "Experienced software developer who builds web applications and writes clear documentation "
    "for recruiting teams reviewing this sample resume."
)


class ParseTests(unittest.TestCase):
    def test_paragraph_docx(self):
        data = _docx_bytes(lambda doc: doc.add_paragraph(f"Alex Example\n{LONG} Skills include Python and SQL."))
        text = extract_text(data, "alex.docx")
        self.assertIn("Python", text)
        self.assertIn("Alex Example", text)

    def test_table_nested_and_header_docx(self):
        def build(doc):
            doc.sections[0].header.paragraphs[0].text = "Header skill HTML"
            doc.sections[0].footer.paragraphs[0].text = "Footer skill CSS"
            doc.add_paragraph(LONG)
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "Power BI"
            table.cell(0, 0).merge(table.cell(0, 1))
            nested = table.cell(0, 0).add_table(rows=1, cols=1)
            nested.cell(0, 0).text = "React"

        text = extract_text(_docx_bytes(build), "table.docx")
        self.assertIn("Power BI", text)
        self.assertEqual(text.lower().count("power bi"), 1)
        self.assertIn("React", text)
        self.assertIn("HTML", text)
        self.assertIn("CSS", text)
        self.assertLess(text.lower().find("html"), text.lower().find("power bi"))

    def test_valid_pdf(self):
        text = extract_text(_pdf(LONG + " Skills include Python and SQL."), "sample.pdf")
        self.assertIn("Python", text)

    def test_empty_corrupt_and_limits(self):
        with self.assertRaises(ValueError) as empty:
            validate_upload("resume.pdf", "application/pdf", 0, 1024)
        self.assertIn("empty", str(empty.exception).lower())

        with self.assertRaises(ValueError) as corrupt_pdf:
            extract_text(b"%PDF-1.4\nnot-a-real-pdf", "bad.pdf")
        self.assertIn("corrupt", str(corrupt_pdf.exception).lower())

        with self.assertRaises(ValueError) as corrupt_docx:
            extract_text(b"PK\x03\x04not-a-zip", "bad.docx")
        self.assertIn("corrupt", str(corrupt_docx.exception).lower())

        oversized = _docx_bytes(lambda doc: doc.add_paragraph(LONG))
        with self.assertRaises(ValueError) as too_big:
            extract_text(oversized, "big.docx", max_docx_uncompressed=50)
        self.assertIn("too large", str(too_big.exception).lower())

        pdf = _pdf(LONG)
        with self.assertRaises(ValueError) as pages:
            extract_text(pdf, "pages.pdf", max_pdf_pages=0)
        self.assertIn("too many pages", str(pages.exception).lower())

        with self.assertRaises(ValueError) as huge_text:
            extract_text(pdf, "long.pdf", max_chars=20)
        self.assertIn("too much text", str(huge_text.exception).lower())

    def test_zip_declared_size_limit(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w") as archive:
            info = zipfile.ZipInfo("word/document.xml")
            payload = b"<w:document>short</w:document>"
            info.file_size = 50_000_000
            archive.writestr(info, payload)
        with self.assertRaises(ValueError) as too_big:
            extract_text(buffer.getvalue(), "bomb.docx", max_docx_uncompressed=1024)
        self.assertTrue("too large" in str(too_big.exception).lower() or "corrupt" in str(too_big.exception).lower())

    def test_password_message(self):
        from app.services.parse import _read_error

        message = str(_read_error(Exception("File has not been decrypted")))
        self.assertIn("password-protected", message)
        self.assertNotIn("OCR is supported", message)


if __name__ == "__main__":
    unittest.main()
